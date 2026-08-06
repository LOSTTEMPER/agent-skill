#!/usr/bin/env python3
"""Build a compact three-page DOCX that demonstrates the full style system."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx.enum.text import WD_BREAK

from format_core import (
    BODY_CJK,
    DEEP_BLUE,
    HEADING_CJK,
    INK,
    LATIN,
    MUTED,
    add_list_paragraph,
    ensure_numbering,
    new_document,
    set_paragraph_shading_and_left_rule,
    set_run_font,
    style_table,
)


def add_formatted_body(doc):
    p = doc.add_paragraph(style="Normal")
    p.add_run("这是一段正文示例，展示微软雅黑、10.5 磅、1.5 倍行距与首行缩进。")
    strong = p.add_run("重点信息使用加粗")
    set_run_font(strong, BODY_CJK, LATIN, 10.5, DEEP_BLUE, bold=True)
    tail = p.add_run("，中英文与数字均使用微软雅黑：")
    set_run_font(tail, BODY_CJK, LATIN, 10.5, INK)
    latin = p.add_run("Product 2026")
    set_run_font(latin, LATIN, LATIN, 10.5, INK)
    end = p.add_run("。")
    set_run_font(end, BODY_CJK, LATIN, 10.5, INK)


def build(output: Path) -> None:
    doc = new_document()
    bullet_id, decimal_id = ensure_numbering(doc)

    title = doc.add_paragraph("文档格式规范示例", style="Title")
    subtitle = doc.add_paragraph("字体 · 行距 · 标题 · 目录 · 列表 · 引用 · 表格", style="Subtitle")
    lead = doc.add_paragraph(
        "本样例只放少量内容，用于快速检查一套中文商务文档的完整视觉层级与独立目录页。",
        style="Lead",
    )

    cover_break = doc.add_paragraph()
    cover_break.add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph("目录", style="Directory Title")
    toc_entries = [
        ("TOC 1", "一、基础排版", "3"),
        ("TOC 2", "1.1　标题层级", "3"),
        ("TOC 3", "1.1.1　三级标题", "3"),
        ("TOC 2", "1.2　列表样式", "3"),
        ("TOC 1", "二、表格与辅助样式", "3"),
        ("TOC 2", "2.1　正文与强调", "3"),
    ]
    for style_name, label, page_number in toc_entries:
        doc.add_paragraph(f"{label}\t{page_number}", style=style_name)
    doc.add_paragraph(
        "样例采用确定性静态目录；正式 Word 文档应在交付前更新目录字段。",
        style="Note",
    )

    toc_break = doc.add_paragraph()
    toc_break.add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph("基础排版", style="Heading 1")
    add_formatted_body(doc)

    doc.add_paragraph("1.1  标题层级", style="Heading 2")
    doc.add_paragraph(
        "二级标题用于主题分组；段前距大于段后距，让标题与后文保持归属。",
        style="Body No Indent",
    )
    doc.add_paragraph("1.1.1  三级标题", style="Heading 3")

    doc.add_paragraph("1.2  列表样式", style="Heading 2")
    add_list_paragraph(doc, "项目符号用于并列信息，换行后与正文文字对齐。", bullet_id)
    add_list_paragraph(doc, "第一步：确认格式要求。", decimal_id)

    quote = doc.add_paragraph(
        "说明框用于关键提示、约束或引用；底色与左侧竖线形成克制强调。",
        style="Quote",
    )
    set_paragraph_shading_and_left_rule(quote)

    doc.add_paragraph("表格与辅助样式", style="Heading 1")
    doc.add_paragraph(
        "表 1　核心样式参数（表题使用 9 磅、居中、灰色）",
        style="Caption",
    )
    table = doc.add_table(rows=4, cols=3)
    data = [
        ("元素", "字号 / 行距", "样式说明"),
        ("正文", "10.5 pt / 1.5", "微软雅黑，首行缩进两字符，段后 6 磅"),
        ("一级标题", "16 pt / 1.15", "微软雅黑加粗、深蓝、段前 18 磅、段后 8 磅"),
        ("表格正文", "9.5 pt / 1.2", "固定列宽、单元格留白、隔行浅色底"),
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(table)

    doc.add_paragraph(
        "注：短字段居中，描述字段左对齐；表头在跨页时重复。",
        style="Note",
    )

    doc.add_paragraph("2.1  正文与强调", style="Heading 2")
    p = doc.add_paragraph(style="Body No Indent")
    r1 = p.add_run("常规文字保持克制；")
    set_run_font(r1, BODY_CJK, LATIN, 10.5, INK)
    r2 = p.add_run("加粗用于结论")
    set_run_font(r2, HEADING_CJK, LATIN, 10.5, DEEP_BLUE, bold=True)
    r3 = p.add_run("，斜体用于术语或作品名。")
    set_run_font(r3, BODY_CJK, LATIN, 10.5, INK, italic=True)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.output.resolve())
    print(f"Created {args.output.resolve()}")


if __name__ == "__main__":
    main()
