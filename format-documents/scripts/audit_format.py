#!/usr/bin/env python3
"""Audit page geometry, named styles, numbering, and fixed table geometry."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from format_core import (
    BODY_CJK,
    MARGIN_BOTTOM_MM,
    MARGIN_SIDE_MM,
    MARGIN_TOP_MM,
    PAGE_HEIGHT_MM,
    PAGE_WIDTH_MM,
    STYLE_TOKENS,
    TABLE_INDENT_DXA,
    TABLE_WIDTH_DXA,
)


def to_mm(length) -> float:
    return length.mm


def close(actual: float, expected: float, tolerance: float = 0.35) -> bool:
    return math.isclose(actual, expected, abs_tol=tolerance)


def audit(path: Path) -> dict:
    doc = Document(path)
    errors: list[str] = []
    warnings: list[str] = []

    for index, section in enumerate(doc.sections, 1):
        checks = {
            "page width": (to_mm(section.page_width), PAGE_WIDTH_MM),
            "page height": (to_mm(section.page_height), PAGE_HEIGHT_MM),
            "top margin": (to_mm(section.top_margin), MARGIN_TOP_MM),
            "bottom margin": (to_mm(section.bottom_margin), MARGIN_BOTTOM_MM),
            "left margin": (to_mm(section.left_margin), MARGIN_SIDE_MM),
            "right margin": (to_mm(section.right_margin), MARGIN_SIDE_MM),
        }
        for label, (actual, expected) in checks.items():
            if not close(actual, expected):
                errors.append(
                    f"Section {index} {label}: expected {expected} mm, got {actual:.2f} mm"
                )
    if doc.settings.odd_and_even_pages_header_footer:
        errors.append("Odd/even page headers must be disabled for consistent page furniture")

    for token in STYLE_TOKENS:
        try:
            style = doc.styles[token.name]
        except KeyError:
            errors.append(f"Missing style: {token.name}")
            continue
        if style.font.size is None or not math.isclose(
            style.font.size.pt, token.size_pt, abs_tol=0.1
        ):
            actual = None if style.font.size is None else style.font.size.pt
            errors.append(f"Style {token.name} size: expected {token.size_pt}, got {actual}")
        rpr = style.element.find(qn("w:rPr"))
        rfonts = None if rpr is None else rpr.find(qn("w:rFonts"))
        if rfonts is None:
            errors.append(f"Style {token.name}: missing explicit rFonts")
        else:
            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                actual_font = rfonts.get(qn(f"w:{slot}"))
                if actual_font != BODY_CJK:
                    errors.append(
                        f"Style {token.name} {slot}: expected {BODY_CJK}, got {actual_font}"
                    )
        spacing = style.paragraph_format.line_spacing
        if isinstance(spacing, float) and not math.isclose(
            spacing, token.line_spacing, abs_tol=0.02
        ):
            errors.append(
                f"Style {token.name} line spacing: expected {token.line_spacing}, got {spacing}"
            )

    numbering = doc.part.numbering_part.element
    formats = {
        node.get(qn("w:val"))
        for node in numbering.findall(".//" + qn("w:numFmt"))
    }
    for required in ("bullet", "decimal", "chineseCounting"):
        if required not in formats:
            errors.append(f"Missing real {required} numbering definition")

    heading_one_style = doc.styles["Heading 1"]
    heading_one_ppr = heading_one_style.element.find(qn("w:pPr"))
    heading_one_num_pr = (
        None if heading_one_ppr is None else heading_one_ppr.find(qn("w:numPr"))
    )
    heading_one_num_id = (
        None if heading_one_num_pr is None else heading_one_num_pr.find(qn("w:numId"))
    )
    heading_one_num_value = (
        None if heading_one_num_id is None else heading_one_num_id.get(qn("w:val"))
    )
    heading_one_definition_ok = False
    if heading_one_num_value is not None:
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) != heading_one_num_value:
                continue
            abstract_ref = num.find(qn("w:abstractNumId"))
            abstract_value = (
                None if abstract_ref is None else abstract_ref.get(qn("w:val"))
            )
            for abstract in numbering.findall(qn("w:abstractNum")):
                if abstract.get(qn("w:abstractNumId")) != abstract_value:
                    continue
                level = abstract.find(qn("w:lvl"))
                num_fmt = None if level is None else level.find(qn("w:numFmt"))
                level_text = None if level is None else level.find(qn("w:lvlText"))
                heading_one_definition_ok = (
                    num_fmt is not None
                    and num_fmt.get(qn("w:val")) == "chineseCounting"
                    and level_text is not None
                    and level_text.get(qn("w:val")) == "%1、"
                )
                break
            break
    if not heading_one_definition_ok:
        errors.append("Heading 1 must use real Chinese numbering: 一、二、三……")

    heading_one_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Heading 1" and paragraph.text.strip()
    ]
    manual_heading_prefix = re.compile(r"^(?:\d+[.、\s]|[一二三四五六七八九十百]+[、.\s])")
    for paragraph in heading_one_paragraphs:
        if manual_heading_prefix.match(paragraph.text.strip()):
            errors.append(
                f"Heading 1 contains a manual number instead of style numbering: {paragraph.text}"
            )

    toc_titles = [p for p in doc.paragraphs if p.style.name == "Directory Title" and p.text.strip()]
    toc_entries = [p for p in doc.paragraphs if p.style.name in {"TOC 1", "TOC 2", "TOC 3"}]
    if len(toc_titles) != 1 or toc_titles[0].text.strip() != "目录":
        errors.append("A single dedicated TOC title paragraph named '目录' is required")
    if not toc_entries:
        errors.append("No TOC 1–3 entries found")
    for paragraph in toc_entries:
        if "\t" not in paragraph.text:
            errors.append(f"TOC entry lacks tab leader/page number separator: {paragraph.text}")
        if paragraph.style.name == "TOC 1":
            label = paragraph.text.split("\t", 1)[0].strip()
            if not re.match(r"^[一二三四五六七八九十百]+、", label):
                errors.append(f"TOC 1 must use Chinese numbering 一、二、三……: {label}")

    page_breaks = sum(
        1
        for br in doc.element.body.iter(qn("w:br"))
        if br.get(qn("w:type")) == "page"
    )
    if page_breaks < 2:
        errors.append(f"Expected at least 2 explicit page breaks, got {page_breaks}")
    if toc_titles:
        paragraphs = list(doc.paragraphs)
        toc_index = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph._p is toc_titles[0]._p
        )
        break_indexes = [
            index
            for index, paragraph in enumerate(paragraphs)
            if any(
                br.get(qn("w:type")) == "page"
                for br in paragraph._p.iter(qn("w:br"))
            )
        ]
        if not any(index < toc_index for index in break_indexes):
            errors.append("TOC page is missing an explicit page break before it")
        if not any(index > toc_index for index in break_indexes):
            errors.append("TOC page is missing an explicit page break after it")

    for t_index, table in enumerate(doc.tables, 1):
        tblpr = table._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        tblind = tblpr.find(qn("w:tblInd"))
        layout = tblpr.find(qn("w:tblLayout"))
        if tblw is None or tblw.get(qn("w:type")) != "dxa":
            errors.append(f"Table {t_index}: missing fixed DXA tblW")
        elif int(tblw.get(qn("w:w"))) != TABLE_WIDTH_DXA:
            errors.append(f"Table {t_index}: tblW is not {TABLE_WIDTH_DXA}")
        if tblind is None or int(tblind.get(qn("w:w"), "0")) != TABLE_INDENT_DXA:
            errors.append(f"Table {t_index}: tblInd is not {TABLE_INDENT_DXA}")
        if layout is None or layout.get(qn("w:type")) != "fixed":
            errors.append(f"Table {t_index}: layout is not fixed")

        grid_widths = [
            int(col.get(qn("w:w")))
            for col in table._tbl.tblGrid.findall(qn("w:gridCol"))
        ]
        if sum(grid_widths) != TABLE_WIDTH_DXA:
            errors.append(
                f"Table {t_index}: grid widths sum to {sum(grid_widths)}, expected {TABLE_WIDTH_DXA}"
            )
        for r_index, row in enumerate(table.rows, 1):
            cell_widths = []
            for cell in row.cells:
                tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                cell_widths.append(int(tcw.get(qn("w:w"))) if tcw is not None else -1)
            if cell_widths != grid_widths:
                errors.append(
                    f"Table {t_index} row {r_index}: tcW {cell_widths} != grid {grid_widths}"
                )

    if not doc.tables:
        warnings.append("No tables found; table styling was not exercised.")

    return {
        "file": str(path),
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "warnings": warnings,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "toc_entries": len(toc_entries),
        "page_breaks": page_breaks,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = audit(args.input.resolve())
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"{result['status']}: {result['file']}")
        for item in result["errors"]:
            print(f"ERROR: {item}")
        for item in result["warnings"]:
            print(f"WARNING: {item}")
    raise SystemExit(0 if result["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
