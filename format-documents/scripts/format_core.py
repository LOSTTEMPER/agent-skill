#!/usr/bin/env python3
"""Shared DOCX formatting primitives for the format-documents skill."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor


BODY_CJK = "Microsoft YaHei"
HEADING_CJK = "Microsoft YaHei"
LATIN = "Microsoft YaHei"

INK = "202124"
DEEP_BLUE = "17365D"
HEADING_BLUE = "365F91"
MUTED = "667085"
PALE_BLUE = "EAF2F8"
TABLE_STRIPE = "F4F7FA"
TABLE_BORDER = "B7C9DA"
WHITE = "FFFFFF"

PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
MARGIN_TOP_MM = 22
MARGIN_BOTTOM_MM = 20
MARGIN_SIDE_MM = 24
HEADER_FOOTER_MM = 12.5
TABLE_WIDTH_DXA = 9184
TABLE_INDENT_DXA = 120
DEFAULT_TABLE_WIDTHS = (1500, 3000, 4684)


@dataclass(frozen=True)
class StyleToken:
    name: str
    cjk_font: str
    latin_font: str
    size_pt: float
    color: str
    bold: bool = False
    italic: bool = False
    alignment: int = WD_ALIGN_PARAGRAPH.LEFT
    before_pt: float = 0
    after_pt: float = 0
    line_spacing: float = 1.0
    first_line_mm: float = 0
    keep_with_next: bool = False


STYLE_TOKENS = (
    StyleToken("Normal", BODY_CJK, LATIN, 10.5, INK, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
               after_pt=6, line_spacing=1.5, first_line_mm=7.4),
    StyleToken("Title", HEADING_CJK, LATIN, 24, DEEP_BLUE, bold=True,
               alignment=WD_ALIGN_PARAGRAPH.CENTER, after_pt=8, line_spacing=1.0,
               keep_with_next=True),
    StyleToken("Subtitle", BODY_CJK, LATIN, 11, MUTED,
               alignment=WD_ALIGN_PARAGRAPH.CENTER, after_pt=16, line_spacing=1.2,
               keep_with_next=True),
    StyleToken("Heading 1", HEADING_CJK, LATIN, 16, DEEP_BLUE, bold=True,
               before_pt=18, after_pt=8, line_spacing=1.15, keep_with_next=True),
    StyleToken("Heading 2", HEADING_CJK, LATIN, 13, HEADING_BLUE, bold=True,
               before_pt=14, after_pt=6, line_spacing=1.15, keep_with_next=True),
    StyleToken("Heading 3", HEADING_CJK, LATIN, 11.5, INK, bold=True,
               before_pt=10, after_pt=4, line_spacing=1.15, keep_with_next=True),
    StyleToken("Body No Indent", BODY_CJK, LATIN, 10.5, INK,
               alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after_pt=6, line_spacing=1.5),
    StyleToken("Lead", BODY_CJK, LATIN, 11.5, DEEP_BLUE,
               after_pt=10, line_spacing=1.35),
    StyleToken("Quote", BODY_CJK, LATIN, 10.5, DEEP_BLUE, italic=True,
               before_pt=8, after_pt=8, line_spacing=1.35),
    StyleToken("Caption", BODY_CJK, LATIN, 9, MUTED,
               alignment=WD_ALIGN_PARAGRAPH.CENTER, before_pt=4, after_pt=6,
               line_spacing=1.15, keep_with_next=True),
    StyleToken("Note", BODY_CJK, LATIN, 8.5, MUTED,
               before_pt=3, after_pt=3, line_spacing=1.2),
    StyleToken("Directory Title", BODY_CJK, LATIN, 18, DEEP_BLUE, bold=True,
               alignment=WD_ALIGN_PARAGRAPH.CENTER, after_pt=20, line_spacing=1.0,
               keep_with_next=True),
    StyleToken("TOC 1", BODY_CJK, LATIN, 11, DEEP_BLUE, bold=True,
               before_pt=8, after_pt=5, line_spacing=1.3),
    StyleToken("TOC 2", BODY_CJK, LATIN, 10.5, INK,
               before_pt=4, after_pt=4, line_spacing=1.3),
    StyleToken("TOC 3", BODY_CJK, LATIN, 10, MUTED,
               before_pt=3, after_pt=3, line_spacing=1.25),
)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, cjk: str, latin: str, size_pt: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    # Use the required CJK-capable face for every slot. Some renderers ignore
    # w:eastAsia and otherwise substitute a Latin-only face.
    run.font.name = cjk
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), cjk)
    rfonts.set(qn("w:hAnsi"), cjk)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), cjk)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def get_or_add_style(doc: DocumentType, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def apply_style_token(doc: DocumentType, token: StyleToken) -> None:
    style = get_or_add_style(doc, token.name)
    style.font.name = token.cjk_font
    style.font.size = Pt(token.size_pt)
    style.font.bold = token.bold
    style.font.italic = token.italic
    style.font.color.rgb = rgb(token.color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), token.cjk_font)
    rfonts.set(qn("w:hAnsi"), token.cjk_font)
    rfonts.set(qn("w:eastAsia"), token.cjk_font)
    rfonts.set(qn("w:cs"), token.cjk_font)
    pf = style.paragraph_format
    pf.alignment = token.alignment
    pf.space_before = Pt(token.before_pt)
    pf.space_after = Pt(token.after_pt)
    pf.line_spacing = token.line_spacing
    pf.first_line_indent = Mm(token.first_line_mm) if token.first_line_mm else None
    pf.keep_with_next = token.keep_with_next
    pf.keep_together = token.keep_with_next
    ppr = style.element.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        ppr.remove(p_bdr)


def configure_toc_styles(doc: DocumentType) -> None:
    for name, indent_mm in (("TOC 1", 0), ("TOC 2", 7), ("TOC 3", 14)):
        style = doc.styles[name]
        pf = style.paragraph_format
        pf.left_indent = Mm(indent_mm) if indent_mm else None
        pf.first_line_indent = None
        tabs = pf.tab_stops
        tabs.clear_all()
        tabs.add_tab_stop(Mm(162), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:shd"))
    if old is not None:
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_margins(cell, top: int = 100, start: int = 120,
                     bottom: int = 100, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:tcMar"))
    if old is not None:
        tcpr.remove(old)
    tc_mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tcpr.append(tc_mar)


def set_paragraph_shading_and_left_rule(paragraph, fill: str = PALE_BLUE,
                                        rule_color: str = DEEP_BLUE) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        ppr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), rule_color)
    p_bdr.append(left)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "120")
    spacing.set(qn("w:after"), "120")
    ppr.append(spacing)


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_field(paragraph, instruction: str, display_text: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, BODY_CJK, LATIN, 8.5, MUTED)


def configure_header_footer(section, label: str = "文档样式规范") -> None:
    header = section.header
    p = header.paragraphs[0]
    _clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run(label)
    set_run_font(left, BODY_CJK, LATIN, 8.5, MUTED)
    right = p.add_run("\t格式规范示例")
    set_run_font(right, BODY_CJK, LATIN, 8.5, MUTED)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Mm(162), WD_TAB_ALIGNMENT.RIGHT)

    footer = section.footer
    fp = footer.paragraphs[0]
    _clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    prefix = fp.add_run("第 ")
    set_run_font(prefix, BODY_CJK, LATIN, 8.5, MUTED)
    add_field(fp, "PAGE", "1")
    suffix = fp.add_run(" 页")
    set_run_font(suffix, BODY_CJK, LATIN, 8.5, MUTED)


def configure_document(doc: DocumentType, header_label: str = "文档样式规范") -> None:
    for token in STYLE_TOKENS:
        apply_style_token(doc, token)
    configure_toc_styles(doc)
    ensure_heading_one_numbering(doc)
    doc.settings.odd_and_even_pages_header_footer = False

    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    for section in doc.sections:
        section.page_width = Mm(PAGE_WIDTH_MM)
        section.page_height = Mm(PAGE_HEIGHT_MM)
        section.top_margin = Mm(MARGIN_TOP_MM)
        section.bottom_margin = Mm(MARGIN_BOTTOM_MM)
        section.left_margin = Mm(MARGIN_SIDE_MM)
        section.right_margin = Mm(MARGIN_SIDE_MM)
        section.header_distance = Mm(HEADER_FOOTER_MM)
        section.footer_distance = Mm(HEADER_FOOTER_MM)
        section.different_first_page_header_footer = False
        configure_header_footer(section, header_label)


def _next_numbering_id(numbering, tag: str) -> int:
    ids = []
    for node in numbering.findall(qn(f"w:{tag}")):
        value = node.get(qn(f"w:{tag}Id"))
        if value is not None and value.isdigit():
            ids.append(int(value))
    return max(ids, default=0) + 1


def _make_abstract_num(abstract_id: int, fmt: str, text_value: str) -> OxmlElement:
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}"[-8:])
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text_value)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "324")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:eastAsia"), HEADING_CJK)
    rpr.append(fonts)
    lvl.append(rpr)
    abstract.append(lvl)
    return abstract


def ensure_numbering(doc: DocumentType) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    ids = []
    for fmt, text_value in (("bullet", "•"), ("decimal", "%1.")):
        abstract_id = _next_numbering_id(numbering, "abstractNum")
        numbering.append(_make_abstract_num(abstract_id, fmt, text_value))
        num_id = _next_numbering_id(numbering, "num")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        ids.append(num_id)
    return ids[0], ids[1]


def ensure_heading_one_numbering(doc: DocumentType) -> int:
    """Attach real Chinese 一、二、三… numbering to the Heading 1 style."""
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "abstractNum")
    abstract = _make_abstract_num(abstract_id, "chineseCounting", "%1、")

    level = abstract.find(qn("w:lvl"))
    if level is None:
        raise RuntimeError("Heading numbering level was not created")

    suffix = level.find(qn("w:suff"))
    if suffix is not None:
        suffix.set(qn("w:val"), "space")

    p_style = OxmlElement("w:pStyle")
    p_style.set(qn("w:val"), "Heading1")
    level.insert(1, p_style)

    level_ppr = level.find(qn("w:pPr"))
    if level_ppr is not None:
        tabs = level_ppr.find(qn("w:tabs"))
        if tabs is not None:
            level_ppr.remove(tabs)
        indent = level_ppr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            level_ppr.append(indent)
        indent.set(qn("w:left"), "0")
        indent.set(qn("w:hanging"), "0")

    numbering.append(abstract)
    num_id = _next_numbering_id(numbering, "num")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)

    heading_style = doc.styles["Heading 1"]
    ppr = heading_style.element.get_or_add_pPr()
    old_num_pr = ppr.find(qn("w:numPr"))
    if old_num_pr is not None:
        ppr.remove(old_num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num_ref))
    ppr.append(num_pr)
    return num_id


def add_list_paragraph(doc: DocumentType, text: str, num_id: int):
    p = doc.add_paragraph(style="Body No Indent")
    p.paragraph_format.left_indent = Mm(12.7)
    p.paragraph_format.first_line_indent = Mm(-6.3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    ppr.append(num_pr)
    run = p.add_run(text)
    set_run_font(run, BODY_CJK, LATIN, 10.5, INK)
    return p


def set_table_geometry(table, widths_dxa: Sequence[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Column widths must sum to {TABLE_WIDTH_DXA} DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblpr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblBorders", "w:tblCellMar"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), TABLE_BORDER)
        borders.append(node)
    tblpr.append(borders)

    tbl_cell_mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", 100), ("start", 120), ("bottom", 100), ("end", 120)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tbl_cell_mar.append(node)
    tblpr.append(tbl_cell_mar)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths_dxa[index]))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Mm(widths_dxa[index] / 1440 * 25.4)
            set_cell_margins(cell)


def style_table(table, widths_dxa: Sequence[int] = DEFAULT_TABLE_WIDTHS) -> None:
    set_table_geometry(table, widths_dxa)
    for r_index, row in enumerate(table.rows):
        row.height = None
        row.height_rule = None
        if r_index == 0:
            trpr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            trpr.append(header)
        for c_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index == 0:
                set_cell_shading(cell, DEEP_BLUE)
            elif r_index % 2 == 0:
                set_cell_shading(cell, TABLE_STRIPE)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                p.style = "Body No Indent"
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.2
                p.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if r_index == 0 or c_index == 1
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in p.runs:
                    set_run_font(
                        run,
                        HEADING_CJK if r_index == 0 else BODY_CJK,
                        LATIN,
                        9.5,
                        WHITE if r_index == 0 else INK,
                        bold=r_index == 0,
                    )


def style_existing_tables(doc: DocumentType) -> None:
    for table in doc.tables:
        column_count = len(table.columns)
        if column_count == 3:
            widths = DEFAULT_TABLE_WIDTHS
        else:
            base, remainder = divmod(TABLE_WIDTH_DXA, column_count)
            widths = tuple(base + (1 if i < remainder else 0) for i in range(column_count))
        style_table(table, widths)


def new_document() -> DocumentType:
    doc = Document()
    configure_document(doc)
    return doc
